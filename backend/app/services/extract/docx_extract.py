from pathlib import Path

from docx import Document

from app.services.extract.base import ExtractionResult, ExtractedSection
from app.services.extract.common import ExtractionError, _headingish, clean_text


class DocxExtractor:
    def extract(self, file_path: str | Path) -> ExtractionResult:
        try:
            doc = Document(file_path)
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Unable to open DOCX: {exc}") from exc

        sections: list[ExtractedSection] = []
        current_section: str | None = None
        buffer: list[str] = []

        def flush() -> None:
            nonlocal buffer
            if buffer:
                content = clean_text("\n".join(buffer))
                if content:
                    sections.append(
                        ExtractedSection(page_number=None, section=current_section, content=content)
                    )
                buffer = []

        def note_heading(text: str) -> None:
            nonlocal current_section
            current_section = text[:150]

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                flush()
                continue
            style = (para.style.name or "").lower()
            if "heading" in style or "title" in style:
                flush()
                note_heading(text)
                buffer.append(text)
            else:
                if _headingish(text) and current_section != text:
                    flush()
                    note_heading(text)
                buffer.append(text)

        # Tables (meeting sign-offs, task lists, etc.)
        for table in doc.tables:
            for row_idx, row in enumerate(table.rows, start=1):
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    buffer.append(f"Table Row {row_idx}: " + " | ".join(cells))
            flush()

        flush()

        if not sections:
            raise ExtractionError("DOCX contained no extractable text.")

        return ExtractionResult(
            text="\n\n".join(s.content for s in sections),
            sections=sections,
            page_count=None,
        )